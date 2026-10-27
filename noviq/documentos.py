# -*- coding: utf-8 -*-
"""📄 Documentos nativos: la IA crea TXT, Markdown, HTML, PDF, Word, Excel, CSV, PowerPoint y JSON en UNA llamada.

Antes, para «hazme un informe en PDF» la IA escribía scripts de Python, instalaba librerías y probaba (muchos pasos
y errores). Ahora la herramienta `crear_documento` ya está lista: la IA solo entrega el título y el contenido
(títulos con #, viñetas con -, listas 1., **negrita** y tablas | a | b |) y el programa genera el archivo con buen
formato. Sin abrir Office ni otros programas.

Librerías: python-docx (Word), openpyxl (Excel), fpdf2 (PDF), python-pptx (PowerPoint); están en requirements.txt."""
import csv
import html
import json
import re
from pathlib import Path

FORMATOS = ["txt", "md", "html", "pdf", "docx", "xlsx", "csv", "pptx", "json"]
_ALIAS = {"word": "docx", "doc": "docx", "excel": "xlsx", "xls": "xlsx", "powerpoint": "pptx", "ppt": "pptx",
          "presentacion": "pptx", "markdown": "md", "texto": "txt", "htm": "html", "web": "html"}


# ───────────── contenido → bloques ─────────────
def bloques(texto):
    """Markdown sencillo → [("h", nivel, txt) | ("li", txt) | ("ol", txt) | ("p", txt) | ("tabla", filas) | ("br",)]."""
    out, tabla = [], []
    for ln in (texto or "").replace("\r\n", "\n").split("\n"):
        s = ln.rstrip()
        if s.strip().startswith("|") and s.strip().endswith("|"):
            celdas = [c.strip() for c in s.strip().strip("|").split("|")]
            if not all(re.fullmatch(r":?-{2,}:?", c) for c in celdas if c):
                tabla.append(celdas)
            continue
        if tabla:
            out.append(("tabla", tabla))
            tabla = []
        m = re.match(r"^(#{1,4})\s+(.*)", s)
        if m:
            out.append(("h", len(m.group(1)), m.group(2).strip()))
        elif re.match(r"^\s*[-*•]\s+", s):
            out.append(("li", re.sub(r"^\s*[-*•]\s+", "", s)))
        elif re.match(r"^\s*\d+[.)]\s+", s):
            out.append(("ol", re.sub(r"^\s*\d+[.)]\s+", "", s)))
        elif s.strip():
            out.append(("p", s.strip()))
        else:
            out.append(("br",))
    if tabla:
        out.append(("tabla", tabla))
    return out


def _plano(s):
    return re.sub(r"\*\*(.+?)\*\*", r"\1", s or "")


def _tablas_extra(tabla):
    """tabla = {"encabezados": [...], "filas": [[...]]} → filas (con encabezado primero)."""
    if not isinstance(tabla, dict):
        return []
    filas = [list(map(str, tabla.get("encabezados") or []))] if tabla.get("encabezados") else []
    return filas + [[("" if c is None else str(c)) for c in f] for f in tabla.get("filas") or []]


# ───────────── generadores ─────────────
def _txt(p, titulo, bl):
    L = ([titulo, "=" * len(titulo), ""] if titulo else [])
    for b in bl:
        if b[0] == "h":
            L += ["", _plano(b[2]).upper() if b[1] == 1 else _plano(b[2]), ""]
        elif b[0] == "li":
            L.append("• " + _plano(b[1]))
        elif b[0] == "ol":
            L.append("- " + _plano(b[1]))
        elif b[0] == "p":
            L.append(_plano(b[1]))
        elif b[0] == "tabla":
            anchos = [max(len(_plano(f[i])) if i < len(f) else 0 for f in b[1]) for i in range(max(len(f) for f in b[1]))]
            for k, f in enumerate(b[1]):
                L.append("  ".join(_plano(f[i] if i < len(f) else "").ljust(anchos[i]) for i in range(len(anchos))))
                if k == 0:
                    L.append("  ".join("-" * a for a in anchos))
        else:
            L.append("")
    p.write_text("\n".join(L).strip() + "\n", encoding="utf-8")


def _md(p, titulo, contenido):
    p.write_text((f"# {titulo}\n\n" if titulo else "") + (contenido or "").strip() + "\n", encoding="utf-8")


def _inline_html(s):
    s = html.escape(s or "")
    s = re.sub(r"\*\*(.+?)\*\*", r"<strong>\1</strong>", s)
    return re.sub(r"\[([^\]]+)\]\((https?://[^)\s]+)\)", r'<a href="\2">\1</a>', s)


def _html(p, titulo, bl):
    cuerpo, lista = [], None
    for b in bl:
        tipo = {"li": "ul", "ol": "ol"}.get(b[0])
        if lista and tipo != lista:
            cuerpo.append(f"</{lista}>")
            lista = None
        if tipo and not lista:
            cuerpo.append(f"<{tipo}>")
            lista = tipo
        if b[0] == "h":
            cuerpo.append(f"<h{b[1] + 1}>{_inline_html(b[2])}</h{b[1] + 1}>")
        elif tipo:
            cuerpo.append(f"<li>{_inline_html(b[1])}</li>")
        elif b[0] == "p":
            cuerpo.append(f"<p>{_inline_html(b[1])}</p>")
        elif b[0] == "tabla":
            filas = b[1]
            cuerpo.append("<table><thead><tr>" + "".join(f"<th>{_inline_html(c)}</th>" for c in filas[0]) + "</tr></thead><tbody>"
                          + "".join("<tr>" + "".join(f"<td>{_inline_html(c)}</td>" for c in f) + "</tr>" for f in filas[1:])
                          + "</tbody></table>")
    if lista:
        cuerpo.append(f"</{lista}>")
    css = ("body{font-family:Segoe UI,Arial,sans-serif;max-width:860px;margin:32px auto;padding:0 18px;line-height:1.6;color:#1d2733}"
           "h1{color:#0f5132;border-bottom:3px solid #3fdc2b;padding-bottom:6px}h2,h3,h4{color:#174a7c;margin-top:1.4em}"
           "table{border-collapse:collapse;width:100%;margin:12px 0}th,td{border:1px solid #c9d3dd;padding:7px 10px;text-align:left}"
           "th{background:#eef4f9}tr:nth-child(even) td{background:#fafcfe}a{color:#174a7c}")
    p.write_text(f"<!DOCTYPE html>\n<html lang=\"es\">\n<head>\n<meta charset=\"UTF-8\">\n"
                 f"<meta name=\"viewport\" content=\"width=device-width, initial-scale=1\">\n<title>{html.escape(titulo or p.stem)}</title>\n"
                 f"<style>{css}</style>\n</head>\n<body>\n" + (f"<h1>{html.escape(titulo)}</h1>\n" if titulo else "")
                 + "\n".join(cuerpo) + "\n</body>\n</html>\n", encoding="utf-8")


def _docx(p, titulo, bl):
    from docx import Document
    from docx.shared import Pt
    doc = Document()
    doc.styles["Normal"].font.name = "Calibri"
    doc.styles["Normal"].font.size = Pt(11)
    if titulo:
        doc.add_heading(titulo, 0)

    def runs(par, s):
        for i, trozo in enumerate(re.split(r"\*\*", s)):
            if trozo:
                par.add_run(trozo).bold = bool(i % 2)
    for b in bl:
        if b[0] == "h":
            doc.add_heading(_plano(b[2]), min(b[1], 4))
        elif b[0] == "li":
            runs(doc.add_paragraph(style="List Bullet"), b[1])
        elif b[0] == "ol":
            runs(doc.add_paragraph(style="List Number"), b[1])
        elif b[0] == "p":
            runs(doc.add_paragraph(), b[1])
        elif b[0] == "tabla":
            filas = b[1]
            n = max(len(f) for f in filas)
            t = doc.add_table(rows=len(filas), cols=n)
            t.style = "Light Grid Accent 1"
            for i, f in enumerate(filas):
                for j in range(n):
                    t.cell(i, j).text = _plano(f[j] if j < len(f) else "")
                    if i == 0:
                        for r in t.cell(i, j).paragraphs[0].runs:
                            r.bold = True
    doc.save(p)


def _numero(v):
    s = str(v).strip()
    if re.fullmatch(r"-?\d+", s):
        return int(s)
    if re.fullmatch(r"-?\d+[.,]\d+", s):
        return float(s.replace(",", "."))
    return v


def _xlsx(p, titulo, hojas):
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter
    wb = Workbook()
    wb.remove(wb.active)
    for nombre, filas in hojas.items():
        ws = wb.create_sheet(str(nombre)[:31] or "Hoja1")
        for f in filas:
            ws.append([_numero(_plano(c)) for c in f])
        if filas:
            for c in ws[1]:
                c.font = Font(bold=True, color="FFFFFF")
                c.fill = PatternFill("solid", fgColor="174A7C")
                c.alignment = Alignment(vertical="center")
            ws.freeze_panes = "A2"
            ws.auto_filter.ref = ws.dimensions
            for i in range(1, ws.max_column + 1):
                ancho = max(len(str(c.value or "")) for c in ws[get_column_letter(i)])
                ws.column_dimensions[get_column_letter(i)].width = min(60, max(10, ancho + 2))
    if not wb.sheetnames:
        wb.create_sheet(titulo[:31] if titulo else "Hoja1")
    wb.save(p)


def _csv(p, filas):
    with open(p, "w", newline="", encoding="utf-8-sig") as f:   # utf-8-sig: Excel abre bien las tildes
        w = csv.writer(f)
        for fila in filas:
            w.writerow([_plano(c) for c in fila])


def _fuente_pdf(pdf):
    """Fuente con tildes y ñ: Arial de Windows o DejaVu; si no hay, la de PDF base (latin-1)."""
    cands = [(Path("C:/Windows/Fonts/arial.ttf"), Path("C:/Windows/Fonts/arialbd.ttf")),
             (Path("/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf"), Path("/usr/share/fonts/truetype/dejavu/DejaVuSans-Bold.ttf")),
             (Path("/System/Library/Fonts/Supplemental/Arial.ttf"), Path("/System/Library/Fonts/Supplemental/Arial Bold.ttf"))]
    for normal, negrita in cands:
        if normal.exists():
            pdf.add_font("base", "", str(normal))
            pdf.add_font("base", "B", str(negrita if negrita.exists() else normal))
            return "base", True
    return "Helvetica", False


def _pdf(p, titulo, bl):
    from fpdf import FPDF
    pdf = FPDF()
    pdf.set_auto_page_break(True, margin=16)
    pdf.add_page()
    fuente, unicode_ok = _fuente_pdf(pdf)

    def t(s):
        s = _plano(s)
        s = re.sub(r"[\U00010000-\U0010FFFF]", "", s)          # sin emojis (las fuentes de PDF no los traen)
        return s if unicode_ok else s.encode("latin-1", "replace").decode("latin-1")
    ancho = pdf.w - pdf.l_margin - pdf.r_margin
    if titulo:
        pdf.set_font(fuente, "B", 18)
        pdf.multi_cell(ancho, 9, t(titulo), new_x="LMARGIN", new_y="NEXT")
        pdf.ln(3)
    for b in bl:
        if b[0] == "h":
            pdf.ln(2)
            pdf.set_font(fuente, "B", {1: 15, 2: 13, 3: 12}.get(b[1], 11))
            pdf.multi_cell(ancho, 7, t(b[2]), new_x="LMARGIN", new_y="NEXT")
            pdf.ln(1)
        elif b[0] in ("li", "ol", "p"):
            pdf.set_font(fuente, "", 11)
            pref = "• " if b[0] == "li" and unicode_ok else ("- " if b[0] in ("li", "ol") else "")
            pdf.multi_cell(ancho, 6, pref + t(b[1]), new_x="LMARGIN", new_y="NEXT")
        elif b[0] == "tabla":
            filas = [[t(c) for c in f] for f in b[1]]
            n = max(len(f) for f in filas)
            pdf.set_font(fuente, "", 9)
            with pdf.table(width=ancho, text_align="LEFT", line_height=5) as tab:
                for f in filas:
                    fila = tab.row()
                    for j in range(n):
                        fila.cell(f[j] if j < len(f) else "")
            pdf.ln(2)
        else:
            pdf.ln(3)
    pdf.output(str(p))


def _pptx(p, titulo, bl, diapositivas=None):
    from pptx import Presentation
    from pptx.util import Pt
    prs = Presentation()
    if not diapositivas:
        diapositivas, actual = [], None
        for b in bl:
            if b[0] == "h" and b[1] <= 2:
                actual = {"titulo": _plano(b[2]), "puntos": []}
                diapositivas.append(actual)
            elif b[0] in ("li", "ol", "p", "h"):
                if actual is None:
                    actual = {"titulo": titulo or "", "puntos": []}
                    diapositivas.append(actual)
                actual["puntos"].append(_plano(b[-1]))
            elif b[0] == "tabla" and actual is not None:
                actual["puntos"] += [" · ".join(map(_plano, f)) for f in b[1]]
    if titulo:
        s = prs.slides.add_slide(prs.slide_layouts[0])
        s.shapes.title.text = titulo
        if len(s.placeholders) > 1:
            s.placeholders[1].text = ""
    for d in diapositivas:
        s = prs.slides.add_slide(prs.slide_layouts[1])
        s.shapes.title.text = str(d.get("titulo") or "")
        cuerpo = s.placeholders[1].text_frame
        cuerpo.clear()
        for i, punto in enumerate((d.get("puntos") or [])[:12]):
            par = cuerpo.paragraphs[0] if i == 0 else cuerpo.add_paragraph()
            par.text = str(punto)
            par.font.size = Pt(20 if len(d.get("puntos") or []) <= 5 else 16)
    prs.save(p)


# ───────────── herramienta ─────────────
def crear(ruta, formato, titulo="", contenido="", tabla=None, hojas=None, diapositivas=None, datos=None):
    """Crea el documento y devuelve la ruta. Lanza ValueError si falta algo."""
    formato = _ALIAS.get(str(formato or "").lower().strip("."), str(formato or "").lower().strip("."))
    p = Path(ruta)
    if not formato:
        formato = p.suffix.lower().strip(".") or "docx"
    if formato not in FORMATOS:
        raise ValueError(f"formato no soportado: {formato}. Usa uno de: {', '.join(FORMATOS)}")
    if p.suffix.lower() != "." + formato:
        p = p.with_suffix("." + formato)
    p.parent.mkdir(parents=True, exist_ok=True)
    bl = bloques(contenido)
    filas_tabla = _tablas_extra(tabla)
    if filas_tabla:
        bl.append(("tabla", filas_tabla))
    if formato == "txt":
        _txt(p, titulo, bl)
    elif formato == "md":
        extra = "\n\n" + "\n".join("| " + " | ".join(f) + " |" + ("\n|" + "---|" * len(f) if i == 0 else "")
                                   for i, f in enumerate(filas_tabla)) if filas_tabla else ""
        _md(p, titulo, (contenido or "") + extra)
    elif formato == "html":
        _html(p, titulo, bl)
    elif formato == "pdf":
        _pdf(p, titulo, bl)
    elif formato == "docx":
        _docx(p, titulo, bl)
    elif formato in ("xlsx", "csv"):
        if isinstance(hojas, dict) and hojas:
            hs = {k: _tablas_extra(v) if isinstance(v, dict) else [list(map(str, f)) for f in v or []] for k, v in hojas.items()}
        else:
            filas = filas_tabla or next((b[1] for b in bl if b[0] == "tabla"), [])
            if not filas:
                filas = [[_plano(b[-1])] for b in bl if b[0] in ("p", "li", "ol", "h")]
            hs = {(titulo or "Hoja1")[:31]: filas}
        if formato == "csv":
            _csv(p, next(iter(hs.values())))
        else:
            _xlsx(p, titulo, hs)
    elif formato == "pptx":
        _pptx(p, titulo, bl, diapositivas if isinstance(diapositivas, list) else None)
    elif formato == "json":
        obj = datos if datos is not None else {"titulo": titulo, "contenido": contenido, "tabla": tabla}
        p.write_text(json.dumps(obj, ensure_ascii=False, indent=2), encoding="utf-8")
    return p


def tool(a, ctx):
    from .oficina import _ruta, respaldar
    formato = a.get("formato") or ""
    titulo = str(a.get("titulo") or "").strip()
    try:
        if a.get("ruta"):
            p = _ruta(ctx, a["ruta"])
        else:   # sin ruta: carpeta documentos/ del trabajo actual (armada aquí: «documentos» NO es tu carpeta Documentos)
            base = re.sub(r"[^\w\- ]+", "", titulo).strip().replace(" ", "_")[:50] or "documento"
            p = ctx.cwd / "documentos" / f"{base}.{_ALIAS.get(formato.lower(), formato.lower()) or 'docx'}"
        bk = respaldar(p) if p.exists() else None
        hecho = crear(p, formato, titulo, a.get("contenido") or "", a.get("tabla"), a.get("hojas"), a.get("diapositivas"), a.get("datos"))
    except ImportError as e:
        return (f"ERROR: falta la librería de Python {e.name}. Reinicia Noviq con iniciar.bat (la instala sola) o ejecuta: "
                "pip install python-docx openpyxl fpdf2 python-pptx")
    except ValueError as e:
        return f"ERROR: {e}"
    except Exception as e:
        return f"ERROR al crear el documento: {type(e).__name__}: {str(e)[:300]}"
    try:
        tam = hecho.stat().st_size
        ctx.emit("file", {"path": str(hecho), "name": hecho.name, "size": tam, "accion": "creado"})
    except Exception:
        tam = 0
    return (f"✔ Documento {hecho.suffix[1:].upper()} creado: {hecho} ({tam // 1024 or 1} KB)"
            + (f" · la versión anterior quedó en {bk}" if bk else "") + ". Dile al usuario dónde quedó; no hace falta revisarlo con otras herramientas.")


TOOL_DEF = {"type": "function", "function": {
    "name": "crear_documento",
    "description": ("Crea un documento LISTO en una sola llamada: txt, md, html, pdf, docx (Word), xlsx (Excel), csv, pptx "
                    "(PowerPoint) o json. Úsala SIEMPRE que el usuario pida un documento, informe, reporte, carta, CV, tabla, hoja "
                    "de cálculo o presentación: NO escribas scripts ni instales librerías para eso. Escribe el contenido con "
                    "# títulos, - viñetas, 1. listas, **negrita** y tablas | col | col |. Excel/CSV: usa 'tabla' u 'hojas'. "
                    "PowerPoint: cada '## título' es una diapositiva (o usa 'diapositivas'). Sin 'ruta' se guarda en la carpeta "
                    "documentos/ del trabajo actual; rutas como 'Escritorio/informe.pdf' van a las carpetas del usuario."),
    "parameters": {"type": "object", "properties": {
        "formato": {"type": "string", "enum": FORMATOS},
        "titulo": {"type": "string"},
        "contenido": {"type": "string", "description": "texto con # títulos, - viñetas, 1. listas, **negrita**, tablas | a | b |"},
        "ruta": {"type": "string", "description": "opcional: carpeta/nombre del archivo"},
        "tabla": {"type": "object", "description": "{\"encabezados\": [...], \"filas\": [[...], ...]}"},
        "hojas": {"type": "object", "description": "xlsx: {\"Ventas\": {\"encabezados\": [...], \"filas\": [[...]]}, ...}"},
        "diapositivas": {"type": "array", "items": {"type": "object"}, "description": "pptx: [{\"titulo\": \"...\", \"puntos\": [\"...\"]}]"},
        "datos": {"description": "json: el objeto a guardar"}},
        "required": ["formato", "contenido"]}}}
