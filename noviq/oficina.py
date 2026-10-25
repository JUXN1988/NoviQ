# -*- coding: utf-8 -*-
"""📎 Agente de oficina sin capturas: Word, Excel, PDF y carpetas con lo que ya hay en el PC.

- Word (.docx) y Excel (.xlsx) se leen y escriben con python-docx / openpyxl (no hace falta abrir Office ni ver la pantalla).
- Convertir a PDF usa el propio Word/Excel/PowerPoint por PowerShell (COM), igual que «Guardar como PDF».
- PDF: leer texto y unir con pypdf.
- Buscar archivos usa el índice de Windows Search (instantáneo) y, si no está, recorre las carpetas del usuario.
- Organizar una carpeta por tipo (primero muestra el plan; con aplicar=true mueve, nunca borra).
Antes de cambiar un documento que ya existe se guarda una copia en .noviq_respaldos junto al archivo."""
import json
import os
import re
import shutil
import subprocess
import time
from pathlib import Path

from .config import IS_WINDOWS, NO_WINDOW

TIPOS = {"Documentos": {".doc", ".docx", ".odt", ".rtf", ".txt", ".md"}, "PDF": {".pdf"},
         "Hojas de calculo": {".xls", ".xlsx", ".xlsm", ".csv", ".ods"}, "Presentaciones": {".ppt", ".pptx", ".odp"},
         "Imagenes": {".jpg", ".jpeg", ".png", ".gif", ".webp", ".bmp", ".heic", ".svg"}, "Videos": {".mp4", ".mkv", ".avi", ".mov", ".wmv"},
         "Musica": {".mp3", ".wav", ".m4a", ".flac", ".ogg"}, "Comprimidos": {".zip", ".rar", ".7z", ".tar", ".gz"},
         "Instaladores": {".exe", ".msi", ".msix"}, "Codigo": {".py", ".js", ".html", ".css", ".json", ".java", ".cs", ".cpp", ".ts"}}


def _ruta(ctx, r):
    p = Path(os.path.expandvars(os.path.expanduser(str(r or "").strip().strip('"'))))
    if not p.is_absolute():
        # «Escritorio/informe.docx», «Documentos\\x.xlsx» → carpetas reales del usuario
        try:
            from .perfil_windows import carpetas_usuario
            cu = carpetas_usuario()
            partes = p.parts
            if partes and partes[0].lower() in {k.lower() for k in cu} | {"desktop", "documents", "downloads"}:
                alias = {"desktop": "Escritorio", "documents": "Documentos", "downloads": "Descargas"}
                clave = next((k for k in cu if k.lower() == alias.get(partes[0].lower(), partes[0]).lower()), None)
                if clave:
                    return cu[clave].joinpath(*partes[1:])
        except Exception:
            pass
        p = ctx.resolve(str(p))
    return p


def respaldar(p):
    """Copia de seguridad antes de cambiar un documento existente."""
    p = Path(p)
    if not p.exists() or p.is_dir():
        return None
    d = p.parent / ".noviq_respaldos"
    d.mkdir(exist_ok=True)
    dst = d / f"{p.stem}_{time.strftime('%Y%m%d_%H%M%S')}{p.suffix}"
    shutil.copy2(p, dst)
    return dst


def _ps(script, timeout=180):
    if not IS_WINDOWS:
        raise RuntimeError("Solo en Windows.")
    r = subprocess.run(["powershell", "-NoProfile", "-ExecutionPolicy", "Bypass", "-Command",
                        "[Console]::OutputEncoding = [System.Text.Encoding]::UTF8\n" + script],
                       capture_output=True, timeout=timeout, creationflags=NO_WINDOW)
    out = r.stdout.decode("utf-8", "replace").strip()
    if r.returncode != 0 and not out:
        raise RuntimeError(r.stderr.decode("utf-8", "replace").strip()[-500:] or "PowerShell falló")
    return out


def _q(s):
    return "'" + str(s).replace("'", "''") + "'"


# ───────────── Word ─────────────
def _runs(par, s):
    for i, trozo in enumerate(re.split(r"\*\*", s)):
        if trozo:
            par.add_run(trozo).bold = bool(i % 2)


def word_crear(p, texto, titulo=None):
    from docx import Document
    doc = Document()
    if titulo:
        doc.add_heading(titulo, 0)
    for ln in (texto or "").splitlines():
        s = ln.rstrip()
        m = re.match(r"^(#{1,3})\s+(.*)", s)
        if m:
            doc.add_heading(m.group(2), len(m.group(1)))
        elif re.match(r"^\s*[-*•]\s+", s):
            _runs(doc.add_paragraph(style="List Bullet"), re.sub(r"^\s*[-*•]\s+", "", s))
        elif re.match(r"^\s*\d+[.)]\s+", s):
            _runs(doc.add_paragraph(style="List Number"), re.sub(r"^\s*\d+[.)]\s+", "", s))
        elif s:
            _runs(doc.add_paragraph(), s)
        else:
            doc.add_paragraph("")
    p.parent.mkdir(parents=True, exist_ok=True)
    doc.save(p)


def word_leer(p, maximo=20000):
    from docx import Document
    doc = Document(p)
    out = [pa.text for pa in doc.paragraphs]
    for t in doc.tables:
        for fila in t.rows:
            out.append(" | ".join(c.text.strip() for c in fila.cells))
    txt = "\n".join(out).strip()
    return txt[:maximo] + ("\n…[recortado]" if len(txt) > maximo else "")


def word_reemplazar(p, buscar, reemplazar):
    from docx import Document
    doc = Document(p)
    n = 0

    def en_parrafos(pars):
        nonlocal n
        for pa in pars:
            if buscar in pa.text:
                texto = pa.text.replace(buscar, reemplazar)
                n += pa.text.count(buscar)
                for i, r in enumerate(pa.runs):
                    r.text = texto if i == 0 else ""
    en_parrafos(doc.paragraphs)
    for t in doc.tables:
        for fila in t.rows:
            for c in fila.cells:
                en_parrafos(c.paragraphs)
    if n:
        doc.save(p)
    return n


# ───────────── Excel ─────────────
def excel_leer(p, hoja=None, rango=None, maximo=200):
    from openpyxl import load_workbook
    wb = load_workbook(p, data_only=True, read_only=True)
    ws = wb[hoja] if hoja else wb.active
    filas = ws[rango] if rango else ws.iter_rows()
    if rango and not isinstance(filas, tuple):
        filas = ((filas,),)
    out = []
    for i, fila in enumerate(filas):
        if i >= maximo:
            out.append(f"…(más filas; hoja con {ws.max_row} filas)")
            break
        vals = ["" if c.value is None else str(c.value) for c in fila]
        if any(vals):
            out.append("\t".join(vals))
    hojas = ", ".join(wb.sheetnames)
    wb.close()
    return f"Hojas: {hojas} · leyendo «{ws.title}»\n" + "\n".join(out)


def excel_escribir(p, hoja=None, celdas=None, filas=None):
    from openpyxl import Workbook, load_workbook
    nuevo = not p.exists()
    wb = Workbook() if nuevo else load_workbook(p)
    if hoja and hoja not in wb.sheetnames:
        ws = wb.create_sheet(hoja) if not (nuevo and wb.active.title == "Sheet") else wb.active
        ws.title = hoja
    else:
        ws = wb[hoja] if hoja else wb.active
    for ref, v in (celdas or {}).items():
        ws[ref] = v
    for fila in filas or []:
        ws.append(list(fila))
    p.parent.mkdir(parents=True, exist_ok=True)
    wb.save(p)
    return ws.title, ws.max_row


# ───────────── PDF ─────────────
def pdf_leer(p, paginas=None, maximo=20000):
    from pypdf import PdfReader
    r = PdfReader(str(p))
    idx = range(len(r.pages))
    if paginas:
        sel = set()
        for parte in str(paginas).split(","):
            if "-" in parte:
                a, b = parte.split("-", 1)
                sel |= set(range(int(a) - 1, int(b)))
            elif parte.strip():
                sel.add(int(parte) - 1)
        idx = [i for i in idx if i in sel]
    txt = "\n\n".join(f"--- página {i + 1} ---\n{(r.pages[i].extract_text() or '').strip()}" for i in idx)
    if not txt.strip() or len(txt) < 30 * max(1, len(list(idx))):
        txt += "\n(Poco texto: puede ser un PDF escaneado. Ábrelo y usa control_pc accion='leer_texto' para leerlo con OCR.)"
    return f"{len(r.pages)} páginas.\n" + txt[:maximo]


def pdf_unir(archivos, salida):
    from pypdf import PdfWriter
    w = PdfWriter()
    for a in archivos:
        w.append(str(a))
    salida.parent.mkdir(parents=True, exist_ok=True)
    with open(salida, "wb") as f:
        w.write(f)


def a_pdf(p, salida=None):
    """Convierte .docx/.doc/.xlsx/.pptx a PDF con el propio Office (COM)."""
    salida = salida or p.with_suffix(".pdf")
    ext = p.suffix.lower()
    if ext in (".doc", ".docx", ".rtf", ".odt"):
        ps = (f"$w = New-Object -ComObject Word.Application; $w.Visible = $false; try {{ $d = $w.Documents.Open({_q(p)}, $false, $true); "
              f"$d.ExportAsFixedFormat({_q(salida)}, 17); $d.Close($false) }} finally {{ $w.Quit() }}")
    elif ext in (".xls", ".xlsx", ".xlsm", ".ods"):
        ps = (f"$x = New-Object -ComObject Excel.Application; $x.Visible = $false; $x.DisplayAlerts = $false; try {{ "
              f"$b = $x.Workbooks.Open({_q(p)}); $b.ExportAsFixedFormat(0, {_q(salida)}); $b.Close($false) }} finally {{ $x.Quit() }}")
    elif ext in (".ppt", ".pptx", ".odp"):
        ps = (f"$pp = New-Object -ComObject PowerPoint.Application; try {{ $pr = $pp.Presentations.Open({_q(p)}, $true, $false, $false); "
              f"$pr.SaveAs({_q(salida)}, 32); $pr.Close() }} finally {{ $pp.Quit() }}")
    else:
        raise RuntimeError("Solo convierto documentos de Word, Excel o PowerPoint.")
    _ps(ps, timeout=240)
    if not Path(salida).exists():
        raise RuntimeError("Office no generó el PDF (¿está instalado y activado?).")
    return salida


# ───────────── Buscar y organizar ─────────────
def buscar(consulta, carpeta=None, maximo=40):
    """Índice de Windows Search (rápido) o recorrido de las carpetas del usuario."""
    q = (consulta or "").strip()
    res = []
    if IS_WINDOWS:
        filtro = q.replace("'", "''")
        ambito = f" AND SCOPE='file:{str(carpeta).replace(chr(39), chr(39) * 2)}'" if carpeta else ""
        sql = (f"SELECT TOP {int(maximo)} System.ItemPathDisplay, System.DateModified FROM SystemIndex WHERE "
               f"(System.FileName LIKE '%{filtro}%' OR CONTAINS('\"{filtro}*\"')){ambito} ORDER BY System.DateModified DESC")
        ps = (f"$c = New-Object -ComObject ADODB.Connection; $r = New-Object -ComObject ADODB.Recordset; "
              f"$c.Open(\"Provider=Search.CollatorDSO;Extended Properties='Application=Windows';\"); $r.Open({_q(sql)}, $c); "
              "while (-not $r.EOF) { $r.Fields.Item('System.ItemPathDisplay').Value; $r.MoveNext() }; $r.Close(); $c.Close()")
        try:
            res = [x for x in _ps(ps, timeout=40).splitlines() if x.strip()]
        except Exception:
            res = []
    if not res:
        from .perfil_windows import carpetas_usuario
        raices = [Path(carpeta)] if carpeta else list(carpetas_usuario().values())
        ql = q.lower()
        vistos = 0
        for raiz in raices:
            for dp, dn, fn in os.walk(raiz):
                dn[:] = [d for d in dn if not d.startswith((".", "node_modules", "$"))]
                for f in fn:
                    vistos += 1
                    if ql in f.lower():
                        res.append(str(Path(dp) / f))
                        if len(res) >= maximo:
                            break
                if len(res) >= maximo or vistos > 200000:
                    break
    return res[:maximo]


def recientes(carpetas, dias=7, maximo=30, extension=None):
    lim = time.time() - dias * 86400
    out = []
    for raiz in carpetas:
        for dp, dn, fn in os.walk(raiz):
            dn[:] = [d for d in dn if not d.startswith((".", "$", "node_modules"))]
            for f in fn:
                if extension and not f.lower().endswith(extension.lower()):
                    continue
                p = Path(dp) / f
                try:
                    m = p.stat().st_mtime
                except OSError:
                    continue
                if m >= lim:
                    out.append((m, p))
    out.sort(reverse=True)
    return [(time.strftime("%Y-%m-%d %H:%M", time.localtime(m)), str(p)) for m, p in out[:maximo]]


def plan_organizar(carpeta):
    plan = []
    for p in sorted(Path(carpeta).iterdir()):
        if p.is_file() and not p.name.startswith((".", "~$", "desktop.ini")):
            cat = next((k for k, exts in TIPOS.items() if p.suffix.lower() in exts), "Otros")
            plan.append((p, Path(carpeta) / cat / p.name))
    return plan


# ───────────── Herramienta para la IA ─────────────
def tool(a, ctx):
    acc = (a.get("accion") or "").lower()
    try:
        if acc == "word_crear":
            p = _ruta(ctx, a.get("ruta") or "Documentos/documento.docx")
            if p.suffix.lower() != ".docx":
                p = p.with_suffix(".docx")
            bk = respaldar(p) if p.exists() else None
            word_crear(p, a.get("texto") or "", a.get("titulo"))
            return f"✔ Documento de Word creado: {p}" + (f" (la versión anterior quedó en {bk})" if bk else "")
        if acc == "word_leer":
            return word_leer(_ruta(ctx, a.get("ruta")))
        if acc == "word_reemplazar":
            p = _ruta(ctx, a.get("ruta"))
            bk = respaldar(p)
            n = word_reemplazar(p, a.get("buscar") or "", a.get("reemplazar") or "")
            return f"✔ {n} reemplazo(s) en {p}" + (f" (respaldo: {bk})" if n and bk else "") if n else "No encontré ese texto en el documento."
        if acc == "excel_leer":
            return excel_leer(_ruta(ctx, a.get("ruta")), a.get("hoja"), a.get("rango"))
        if acc == "excel_escribir":
            p = _ruta(ctx, a.get("ruta") or "Documentos/libro.xlsx")
            bk = respaldar(p) if p.exists() else None
            celdas = a.get("celdas")
            if isinstance(celdas, str):
                celdas = json.loads(celdas)
            hoja, n = excel_escribir(p, a.get("hoja"), celdas, a.get("filas"))
            return f"✔ Excel guardado: {p} (hoja «{hoja}», {n} filas)" + (f" · respaldo: {bk}" if bk else "")
        if acc == "pdf_leer":
            return pdf_leer(_ruta(ctx, a.get("ruta")), a.get("paginas"))
        if acc == "pdf_unir":
            archivos = [_ruta(ctx, x) for x in a.get("archivos") or []]
            if len(archivos) < 2:
                return "ERROR: pasa al menos 2 PDF en 'archivos'."
            sal = _ruta(ctx, a.get("salida") or "Documentos/unido.pdf")
            pdf_unir(archivos, sal)
            return f"✔ PDF unido: {sal} ({len(archivos)} archivos)"
        if acc == "a_pdf":
            p = _ruta(ctx, a.get("ruta"))
            return f"✔ PDF creado: {a_pdf(p, _ruta(ctx, a['salida']) if a.get('salida') else None)}"
        if acc == "buscar":
            res = buscar(a.get("consulta") or a.get("texto") or "", _ruta(ctx, a["carpeta"]) if a.get("carpeta") else None)
            return ("Encontrados:\n" + "\n".join(res)) if res else "No encontré archivos con ese nombre o contenido."
        if acc == "recientes":
            from .perfil_windows import carpetas_usuario
            cs = [_ruta(ctx, a["carpeta"])] if a.get("carpeta") else list(carpetas_usuario().values())[:3]
            res = recientes(cs, int(a.get("dias") or 7), extension=a.get("extension"))
            return ("Archivos modificados recientemente:\n" + "\n".join(f"{t}  {p}" for t, p in res)) if res else "No hay archivos recientes."
        if acc == "organizar":
            c = _ruta(ctx, a.get("carpeta") or "Descargas")
            plan = plan_organizar(c)
            if not plan:
                return f"No hay archivos sueltos que ordenar en {c}."
            if not a.get("aplicar"):
                resumen = {}
                for _, dst in plan:
                    resumen[dst.parent.name] = resumen.get(dst.parent.name, 0) + 1
                return (f"Plan para ordenar {c} ({len(plan)} archivos, no se borra nada):\n"
                        + "\n".join(f"- {k}: {v} archivos" for k, v in sorted(resumen.items()))
                        + "\nPara hacerlo llama de nuevo con aplicar=true.")
            movidos = 0
            for src, dst in plan:
                dst.parent.mkdir(exist_ok=True)
                if dst.exists():
                    dst = dst.with_name(f"{dst.stem}_{int(time.time())}{dst.suffix}")
                shutil.move(str(src), str(dst))
                movidos += 1
            return f"✔ Ordené {movidos} archivos de {c} en carpetas por tipo (no se borró nada)."
        return ("ERROR: accion debe ser word_crear, word_leer, word_reemplazar, excel_leer, excel_escribir, pdf_leer, pdf_unir, "
                "a_pdf, buscar, recientes u organizar.")
    except ImportError as e:
        return f"ERROR: falta una librería de Python ({e.name}). Instálala con: pip install python-docx openpyxl pypdf"
    except FileNotFoundError as e:
        return f"ERROR: no existe {e.filename or ''}"
    except Exception as e:
        return f"ERROR: {type(e).__name__}: {str(e)[:400]}"


TOOL_DEF = {"type": "function", "function": {
    "name": "oficina",
    "description": "Word, Excel, PDF y carpetas SIN abrir programas ni mirar la pantalla (rápido y barato): crear/leer/reemplazar en Word, leer/escribir Excel, leer/unir PDF, convertir a PDF con Office, buscar archivos con el índice de Windows, archivos recientes y ordenar una carpeta por tipo. Rutas como 'Escritorio/informe.docx' o 'Descargas' usan las carpetas reales del usuario.",
    "parameters": {"type": "object", "properties": {
        "accion": {"type": "string", "enum": ["word_crear", "word_leer", "word_reemplazar", "excel_leer", "excel_escribir", "pdf_leer",
                                              "pdf_unir", "a_pdf", "buscar", "recientes", "organizar"]},
        "ruta": {"type": "string"}, "titulo": {"type": "string"},
        "texto": {"type": "string", "description": "word_crear: contenido (# títulos, - viñetas, **negrita**)"},
        "buscar": {"type": "string"}, "reemplazar": {"type": "string"},
        "hoja": {"type": "string"}, "rango": {"type": "string", "description": "ej. A1:D20"},
        "celdas": {"type": "object", "description": "excel_escribir: {\"A1\": \"Total\", \"B1\": 120}"},
        "filas": {"type": "array", "items": {"type": "array", "items": {}}, "description": "excel_escribir: filas a agregar al final"},
        "archivos": {"type": "array", "items": {"type": "string"}}, "salida": {"type": "string"}, "paginas": {"type": "string"},
        "consulta": {"type": "string"}, "carpeta": {"type": "string"}, "dias": {"type": "integer"}, "extension": {"type": "string"},
        "aplicar": {"type": "boolean"}}, "required": ["accion"]}}}
