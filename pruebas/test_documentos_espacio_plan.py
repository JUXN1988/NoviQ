# -*- coding: utf-8 -*-
"""Pruebas de 📄 documentos nativos, 🗂 carpetas propias / proyecto automático / preparador, y ⏹ cancelar el plan.
    python -m pytest -q pruebas/test_documentos_espacio_plan.py"""
import os
import sys
import tempfile
import threading
import time
from pathlib import Path

os.environ.setdefault("NOVIQ_DATOS", tempfile.mkdtemp(prefix="noviq_test_"))
sys.path.insert(0, str(Path(__file__).resolve().parents[1]))

import pytest  # noqa: E402

from noviq import documentos as DOC  # noqa: E402
from noviq import espacio as ES  # noqa: E402

CONTENIDO = """# Informe de ventas
Resumen **importante** del mes en Santo Domingo: ñandú, acción, café.
## Detalle
- Producto A
- Producto B
1. Primero
| Producto | Unidades | Precio |
|---|---|---|
| Café | 10 | 2,50 |
| Té | 4 | 1.75 |
"""


@pytest.mark.parametrize("fmt", DOC.FORMATOS)
def test_crea_todos_los_formatos(tmp_path, fmt):
    p = DOC.crear(tmp_path / "doc", fmt, "Ventas de octubre", CONTENIDO)
    assert p.exists() and p.stat().st_size > 0 and p.suffix == "." + fmt
    if fmt == "docx":
        from docx import Document
        d = Document(p)
        assert "Ventas de octubre" in [x.text for x in d.paragraphs] and d.tables[0].cell(1, 0).text == "Café"
    elif fmt == "xlsx":
        from openpyxl import load_workbook
        ws = load_workbook(p).active
        assert ws["A1"].value == "Producto" and ws["B2"].value == 10 and ws["C3"].value == 1.75 and ws["A1"].font.bold
    elif fmt == "pdf":
        from pypdf import PdfReader
        txt = PdfReader(str(p)).pages[0].extract_text()
        assert "Ventas de octubre" in txt and "acción" in txt
    elif fmt == "pptx":
        from pptx import Presentation
        assert len(Presentation(p).slides) >= 2
    elif fmt == "html":
        h = p.read_text(encoding="utf-8")
        assert "<table>" in h and "<strong>importante</strong>" in h and "<ul>" in h
    elif fmt == "csv":
        assert p.read_text(encoding="utf-8-sig").splitlines()[1].startswith("Café,10")


def test_excel_con_varias_hojas(tmp_path):
    p = DOC.crear(tmp_path / "libro.xlsx", "excel", "x", "", hojas={"Ventas": {"encabezados": ["A", "B"], "filas": [[1, 2]]},
                                                                     "Gastos": {"encabezados": ["C"], "filas": [[3]]}})
    from openpyxl import load_workbook
    assert load_workbook(p).sheetnames == ["Ventas", "Gastos"]


def test_herramienta_y_permisos(tmp_path):
    from noviq import herr_windows as HW
    from noviq import tools as T
    assert "crear_documento" in T.FUNCS and any(d["function"]["name"] == "crear_documento" for d in T.BUILTIN)
    assert HW.necesita_permiso("crear_documento", {"formato": "pdf"}, "preguntar") is False
    assert HW.necesita_permiso("crear_documento", {"ruta": "Escritorio/x.pdf"}, "preguntar") is True
    assert HW.necesita_permiso("crear_documento", {"ruta": "Escritorio/x.pdf"}, "auto") is False
    conv = {"id": "doc1", "project": None, "carpeta": str(tmp_path / "trabajo"), "messages": []}
    ev = []
    ctx = T.ToolContext(conv, lambda k, p: ev.append(k), threading.Event())
    r = DOC.tool({"formato": "pdf", "titulo": "Noticias de hoy", "contenido": "- uno\n- dos"}, ctx)
    assert r.startswith("✔") and (tmp_path / "trabajo" / "documentos" / "Noticias_de_hoy.pdf").exists() and "file" in ev


def test_nivel_documento():
    from noviq import via_rapida as VRP
    for t in ("Hazme un informe en PDF sobre el clima de Quito", "redacta una carta de renuncia en Word",
              "crea una presentación de 5 diapositivas sobre reciclaje", "pásame esta tabla a Excel"):
        assert VRP.nivel(t) == "documento", t
    assert VRP.nivel("haz un script en Python que genere un PDF") != "documento"
    assert VRP.nivel("crea una página web de recetas") == "grande"


# ───────────── 🗂 carpetas ─────────────
def test_conversacion_nueva_tiene_carpeta_propia():
    from noviq import tools as T
    from noviq.config import general_dir
    from noviq.server_state import new_conv
    c = new_conv("abcd1234")
    c["title"] = "Noticias de Ecuador"
    cwd = T.ToolContext(c, lambda *a: None, threading.Event()).cwd
    assert cwd.parent == general_dir() / "conversaciones" and "noticias_de_ecuador" in cwd.name and cwd.exists()
    assert T.ToolContext(c, lambda *a: None, threading.Event()).cwd == cwd          # estable
    vieja = {"id": "old", "project": None, "messages": []}                            # de antes: sigue en _General
    assert T.ToolContext(vieja, lambda *a: None, threading.Event()).cwd == general_dir()
    c2 = new_conv("ef567890", project="_General")                                     # «_General» elegido en la ventana
    assert T.ToolContext(c2, lambda *a: None, threading.Event()).cwd.parent == general_dir() / "conversaciones"


def test_prompt_de_conversacion_nueva_no_ve_lo_de_general(monkeypatch):
    from noviq import agent as AG
    from noviq.config import general_dir
    from noviq.server_state import new_conv
    (general_dir() / "config_viejo.json").write_text("{}", encoding="utf-8")
    c = new_conv("gh901234", project="_General")
    c["messages"] = [{"role": "user", "content": "explícame la recursividad"}]
    s = AG.system_prompt(c)
    assert "config_viejo.json" not in s and "Archivos del proyecto" not in s


def test_nombre_de_proyecto_unico():
    from noviq.config import projects_root
    n1 = ES.nombre_proyecto("Crea una página web de recetas con buscador")
    assert n1.lower().startswith("pagina web") or n1.lower().startswith("página web")
    (projects_root() / n1).mkdir(parents=True, exist_ok=True)
    assert ES.nombre_proyecto("Crea una página web de recetas con buscador") == n1 + " 2"
    assert ES.nombre_proyecto("x", "Mi Tienda") .startswith("Mi Tienda")


def test_programa_nuevo_crea_su_proyecto(monkeypatch):
    from noviq import agent as AG
    from noviq import cerebro as CB
    from noviq.server_state import new_conv
    monkeypatch.setattr(AG, "stream_chat", lambda *a, **k: ("Listo.", "", [], "stop"))
    monkeypatch.setattr(CB, "plan", lambda *a, **k: False)
    monkeypatch.setattr(AG, "fallback_chain", lambda conv, first=None: ["prueba::modelo"])
    monkeypatch.setattr(AG.EV, "modelo_inicial", lambda conv, want: ("prueba::modelo", None))
    monkeypatch.setattr(AG.EV, "registrar", lambda *a, **k: None)
    c = new_conv("ij23456789")
    c["messages"] = [{"role": "user", "content": "Crea una calculadora de propinas en HTML\n\n(Es un PROYECTO NUEVO: créalo primero "
                                                 "con crear_proyecto con un nombre corto y claro y trabaja dentro de él.)"}]
    ev = []
    AG.run_agent(c, lambda k, p: ev.append((k, p)), threading.Event())
    assert ES.es_real(c["project"]) and "calculadora" in c["project"].lower()
    assert "ya fue creado por Noviq" in c["messages"][0]["content"] and "créalo primero" not in c["messages"][0]["content"]
    assert any(k == "project" for k, _ in ev)


def test_preparador_crea_base_sin_salirse(tmp_path):
    plan = {"resumen": "Calculadora de propinas", "tecnologia": "HTML, CSS y JS",
            "estructura": ["index.html: página principal", "css/estilos.css: estilos", "js/app.js: lógica", "img/",
                           "../fuera.js: no debe crearse", "datos.bin: binario"],
            "tareas": ["Maquetar", "Lógica", "Probar"], "pruebas": "abrir index.html"}
    creados = ES.preparar(tmp_path, plan)
    assert {"index.html", "css/estilos.css", "js/app.js", "img/", "PLAN.md", "README.md"} <= set(creados)
    assert not (tmp_path.parent / "fuera.js").exists() and not (tmp_path / "datos.bin").exists()
    html = (tmp_path / "index.html").read_text(encoding="utf-8")
    assert 'href="css/estilos.css"' in html and 'src="js/app.js"' in html and "<!DOCTYPE html>" in html
    assert "1. Maquetar" in (tmp_path / "PLAN.md").read_text(encoding="utf-8")
    (tmp_path / "js" / "app.js").write_text("// mío", encoding="utf-8")
    ES.preparar(tmp_path, plan)                                         # no pisa lo existente
    assert (tmp_path / "js" / "app.js").read_text(encoding="utf-8") == "// mío"


# ───────────── ⏹ cancelar el plan ─────────────
def _revision_en_hilo(conv, stop):
    from noviq import cerebro as CB
    from noviq.config import CONFIG
    CONFIG["revisar_plan_seg"] = 30
    res, evs = {}, []
    th = threading.Thread(target=lambda: res.setdefault("plan", CB.review(conv, lambda k, p: evs.append((k, p)), stop,
                                                                          {"tareas": ["a", "b", "c"]}, "x::m")))
    th.start()
    t0 = time.time()
    while not CB.pendiente(conv["id"]) and time.time() - t0 < 3:
        time.sleep(0.02)
    return th, res, evs


@pytest.mark.parametrize("texto", ["no", "cancela", "No, cancélalo", "detente", "mejor no"])
def test_escribir_no_en_el_chat_detiene_el_trabajo(texto):
    from noviq import cerebro as CB
    stop = threading.Event()
    th, res, evs = _revision_en_hilo({"id": "pl1"}, stop)
    assert CB.respuesta_chat("pl1", texto) == "detenido"
    th.join(3)
    assert res["plan"] is None and stop.is_set() and any("detuve el trabajo" in p.get("text", "") for k, p in evs if k == "notice")


def test_si_aprueba_y_otro_texto_es_cambio():
    from noviq import cerebro as CB
    stop = threading.Event()
    th, res, _ = _revision_en_hilo({"id": "pl2"}, stop)
    assert CB.respuesta_chat("pl2", "sí, dale") == "aprobado"
    th.join(3)
    assert res["plan"]["tareas"] == ["a", "b", "c"] and not stop.is_set()
    th, res, _ = _revision_en_hilo({"id": "pl3"}, stop)
    assert CB.respuesta_chat("pl3", "agrega también un modo oscuro") == "cambio"
    th.join(3)
    assert res["plan"] is not None
    assert CB.respuesta_chat("sin_plan_pendiente", "no") is None


def test_boton_cancelar_detiene_y_sin_plan_sigue():
    from noviq import cerebro as CB
    stop = threading.Event()
    th, res, _ = _revision_en_hilo({"id": "pl4"}, stop)
    CB.responder(CB.pendiente("pl4"), "detener")
    th.join(3)
    assert res["plan"] is None and stop.is_set()
    stop2 = threading.Event()
    th, res, _ = _revision_en_hilo({"id": "pl5"}, stop2)
    CB.responder(CB.pendiente("pl5"), "sin_plan")
    th.join(3)
    assert res["plan"] is None and not stop2.is_set()
